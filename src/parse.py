from openpyxl import load_workbook
from datetime import datetime
from dataclasses import dataclass
from typing import cast, Optional
import re


@dataclass
class Person:
    name: str
    vorname: str
    untersuchung: Optional[datetime]
    unterweisung: Optional[datetime]
    anlage: Optional[datetime]
    agtUebung: Optional[datetime]


@dataclass
class ParsedData:
    people: list[Person]
    data_age: Optional[str]


def parse(xlsx: str) -> ParsedData:
    wb = load_workbook(xlsx)
    ws = wb.active
    if not ws:
        print("Empty xlsx sheet")
        return ParsedData(people=[], data_age=None)

    data_age_val = ws["A2"].value
    data_age: Optional[str] = None
    if data_age_val:
        if isinstance(data_age_val, str):
            first_line = data_age_val.split("\n")[0]
            r_paren_index = first_line.rfind(")")
            if r_paren_index != -1:
                data_age = first_line[: r_paren_index + 1].strip()
            else:
                data_age = first_line.strip()
        elif isinstance(data_age_val, datetime):
            data_age = data_age_val.strftime(
                "%d.%m.%Y"
            )  # Convert datetime to string if it's a datetime object

    people = []

    for row in ws.iter_rows(min_row=5, max_col=12):
        (
            _,
            status,
            _,
            name,
            vorname,
            untersuchung_val,
            unterweisung_val,
            anlage_val,
            agtUebung_val,
            _,
            _,
            _,
        ) = (c.value for c in row)

        if status != "Aktiv":
            continue

        people.append(
            Person(
                str(name),
                str(vorname),
                cast(datetime, untersuchung_val) if untersuchung_val else None,
                cast(datetime, unterweisung_val) if unterweisung_val else None,
                cast(datetime, anlage_val) if anlage_val else None,
                cast(datetime, agtUebung_val) if agtUebung_val else None,
            )
        )

    return ParsedData(people=people, data_age=data_age)
